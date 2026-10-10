"""
Local fallback for querying tabular resources when no Tabular API is available.

Downloads the resource file (CSV/TSV/XLSX, size-capped, SSRF-hardened), then
filters, sorts and paginates it in-process. Results use the same shape as the
Tabular API (``data`` / ``meta`` / ``links``) so tools can format both alike.
"""

import asyncio
import csv
import inspect
import io
import logging
import os
import time
import unicodedata
import zipfile
from collections import OrderedDict
from collections.abc import AsyncIterator, Callable, Iterator
from datetime import date, datetime
from typing import Any
from urllib.parse import urlsplit

from helpers import file_cache
from helpers.logging import MAIN_LOGGER_NAME
from helpers.ssrf import ssrf_async_session

logger = logging.getLogger(MAIN_LOGGER_NAME)

DEFAULT_MAX_BYTES = 25 * 1024 * 1024
CACHE_TTL_SECONDS = 600
CACHE_MAX_ENTRIES = 4
_CHUNK_SIZE = 64 * 1024
_CSV_DELIMITERS = ",;\t|"

SUPPORTED_FORMATS = ("csv", "tsv", "xlsx")

# resource_id -> (url, version, fetched_at, content)
_file_cache: OrderedDict[str, tuple[str, str, float, bytes]] = OrderedDict()


class LocalQueryError(Exception):
    """Raised when a resource cannot be queried locally (message is LLM-facing)."""


def max_bytes() -> int:
    """
    Maximum file size for local queries, from LOCAL_QUERY_MAX_BYTES.

    Defaults to 25 MB. 0 disables the local fallback.
    """
    raw = os.getenv("LOCAL_QUERY_MAX_BYTES", "").strip()
    if not raw:
        return DEFAULT_MAX_BYTES
    try:
        return max(0, int(raw))
    except ValueError:
        return DEFAULT_MAX_BYTES


def is_enabled() -> bool:
    return max_bytes() > 0


def format_size(size: int) -> str:
    if size < 1024 * 1024:
        return f"{size / 1024:.0f} KB"
    return f"{size / (1024 * 1024):.0f} MB"


def detect_format(resource: dict[str, Any]) -> str | None:
    """Return csv|tsv|xlsx for a resource payload, or None if unsupported."""
    candidates = [
        str(resource.get("format") or "").lower().lstrip("."),
        os.path.splitext(urlsplit(str(resource.get("url") or "")).path)[1]
        .lower()
        .lstrip("."),
    ]
    mime = str(resource.get("mime") or "").lower()
    if "spreadsheetml" in mime:
        candidates.append("xlsx")
    elif mime == "text/tab-separated-values":
        candidates.append("tsv")
    elif mime == "text/csv":
        candidates.append("csv")
    for candidate in candidates:
        if candidate in SUPPORTED_FORMATS:
            return candidate
    return None


def clear_cache() -> None:
    """Clear the in-memory downloaded files cache. Useful for testing."""
    _file_cache.clear()


def _remember(resource_id: str, url: str, version: str, content: bytes) -> bytes:
    _file_cache[resource_id] = (url, version, time.time(), content)
    _file_cache.move_to_end(resource_id)
    while len(_file_cache) > CACHE_MAX_ENTRIES:
        _file_cache.popitem(last=False)
    return content


def _resource_version(resource: dict[str, Any]) -> str:
    """What the catalog knows about the file's current version (may be empty)."""
    checksum = resource.get("checksum")
    checksum_value = checksum.get("value") if isinstance(checksum, dict) else None
    return "|".join(
        str(part) for part in (resource.get("last_modified"), checksum_value) if part
    )


async def iter_chunks(resp: Any) -> AsyncIterator[bytes]:
    """Yield body chunks from a streamed response (async, or sync when mocked)."""
    stream = resp.iter_content(_CHUNK_SIZE)
    if inspect.isawaitable(stream):
        stream = await stream
    if hasattr(stream, "__aiter__"):
        async for chunk in stream:
            yield chunk
    else:
        for chunk in stream:
            yield chunk


async def _download(resource_id: str, url: str, version: str = "") -> bytes:
    """
    Get a resource file: from memory, from the on-disk cache, or from its URL.

    `version` is what the catalog says about the file (see _resource_version).
    A cached copy is used as is while the version matches and the copy was
    checked recently; after that the source is asked whether it changed
    (conditional request), and the file is only downloaded again if it did.
    """
    limit = max_bytes()
    cached = _file_cache.get(resource_id)
    if (
        cached
        and cached[:2] == (url, version)
        and time.time() - cached[2] < CACHE_TTL_SECONDS
    ):
        _file_cache.move_to_end(resource_id)
        if file_cache.is_enabled():
            file_cache.touch(url)
        return cached[3]

    headers: dict[str, str] = {}
    on_disk = (
        await asyncio.to_thread(file_cache.lookup, url)
        if file_cache.is_enabled()
        else None
    )
    if on_disk and on_disk.get("version", "") == version:
        age = time.time() - float(on_disk.get("checked_at") or 0)
        if age < file_cache.revalidate_seconds():
            content = await asyncio.to_thread(file_cache.read, url)
            if content is not None:
                logger.info(f"Local query: resource {resource_id} read from disk cache")
                return _remember(resource_id, url, version, content)
        if on_disk.get("etag"):
            headers["If-None-Match"] = on_disk["etag"]
        if on_disk.get("http_last_modified"):
            headers["If-Modified-Since"] = on_disk["http_last_modified"]

    too_large = LocalQueryError(
        f"The file is larger than the {format_size(limit)} limit for local queries. "
        "Fetch the resource URL directly instead."
    )
    session = ssrf_async_session()
    try:
        logger.info(f"Local query: downloading resource {resource_id} from {url}")
        resp = await session.get(
            url,
            headers=headers or None,
            timeout=60.0,
            allow_redirects=True,
            stream=True,
        )
        not_modified = bool(headers) and resp.status_code == 304
        chunks: list[bytes] = []
        if not not_modified:
            resp.raise_for_status()
            declared = resp.headers.get("content-length")
            if declared and declared.isdigit() and int(declared) > limit:
                raise too_large
            received = 0
            async for chunk in iter_chunks(resp):
                received += len(chunk)
                if received > limit:
                    raise too_large
                chunks.append(chunk)
        etag = resp.headers.get("etag")
        http_last_modified = resp.headers.get("last-modified")
    finally:
        await session.close()

    if not_modified:
        content = await asyncio.to_thread(file_cache.read, url)
        if content is None:
            # The cached file disappeared in the meantime: download it in full
            await asyncio.to_thread(file_cache.remove, url)
            return await _download(resource_id, url, version)
        logger.info(f"Local query: resource {resource_id} unchanged at the source")
        await asyncio.to_thread(file_cache.mark_checked, url)
        return _remember(resource_id, url, version, content)

    content = b"".join(chunks)
    if file_cache.is_enabled():
        await asyncio.to_thread(
            file_cache.store,
            url,
            content,
            resource_id=resource_id,
            version=version,
            etag=etag,
            http_last_modified=http_last_modified,
        )
    return _remember(resource_id, url, version, content)


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("latin-1")


def _sniff_delimiter(text: str, file_format: str) -> str:
    if file_format == "tsv":
        return "\t"
    sample = text[:_CHUNK_SIZE]
    try:
        return csv.Sniffer().sniff(sample, delimiters=_CSV_DELIMITERS).delimiter
    except csv.Error:
        header = sample.split("\n", 1)[0]
        return max(_CSV_DELIMITERS, key=header.count)


def _cell_to_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        if value.time() == datetime.min.time():
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _iter_csv(content: bytes, file_format: str) -> Iterator[list[str]]:
    text = _decode(content)
    reader = csv.reader(
        io.StringIO(text, newline=""), delimiter=_sniff_delimiter(text, file_format)
    )
    yield from reader


def _iter_xlsx(content: bytes) -> Iterator[list[str]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError) as e:
        raise LocalQueryError(f"The file is not a valid XLSX workbook ({e}).") from e
    try:
        sheet = workbook.worksheets[0]
        for row in sheet.iter_rows(values_only=True):
            yield [_cell_to_str(cell) for cell in row]
    finally:
        workbook.close()


def _column_names(header: list[str]) -> list[str]:
    names: list[str] = []
    seen: set[str] = set()
    for index, raw in enumerate(header, 1):
        name = raw.strip() or f"column_{index}"
        base, suffix = name, 2
        while name in seen:
            name = f"{base}_{suffix}"
            suffix += 1
        seen.add(name)
        names.append(name)
    return names


def _to_number(value: str) -> float | None:
    text = value.strip().replace(" ", "").replace(" ", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        pass
    # Decimal comma ("1234,56"), common in Portuguese and French files
    if text.count(",") == 1 and "." not in text:
        try:
            return float(text.replace(",", "."))
        except ValueError:
            pass
    return None


def _build_filter(operator: str, expected: str) -> Callable[[str], bool]:
    """Same semantics as the Tabular API: less/greater are inclusive."""
    if operator == "contains":
        needle = expected.casefold()
        return lambda cell: needle in cell.casefold()

    expected_number = _to_number(expected)

    def compare(cell: str) -> int | None:
        """-1, 0 or 1 for cell vs expected; None when the cell is empty."""
        if not cell.strip():
            return None
        if expected_number is not None:
            cell_number = _to_number(cell)
            if cell_number is not None:
                return (cell_number > expected_number) - (cell_number < expected_number)
        return (cell > expected) - (cell < expected)

    checks: dict[str, Callable[[int], bool]] = {
        "exact": lambda order: order == 0,
        "less": lambda order: order <= 0,
        "greater": lambda order: order >= 0,
        "strictly_less": lambda order: order < 0,
        "strictly_greater": lambda order: order > 0,
    }
    check = checks[operator]

    if operator == "exact":
        return lambda cell: cell == expected or compare(cell) == 0

    def matches(cell: str) -> bool:
        order = compare(cell)
        return order is not None and check(order)

    return matches


def _text_sort_key(value: str) -> str:
    """Case- and accent-insensitive key, so "Évora" sorts next to "Evora"."""
    decomposed = unicodedata.normalize("NFKD", value.casefold())
    return "".join(char for char in decomposed if not unicodedata.combining(char))


def _sort_rows(rows: list[list[str]], index: int, descending: bool) -> None:
    """Sort in place: numerically if every non-empty value is a number; blanks last."""
    filled = [row for row in rows if row[index].strip()]
    blank = [row for row in rows if not row[index].strip()]
    numbers = [_to_number(row[index]) for row in filled]
    if all(number is not None for number in numbers):
        order = sorted(
            range(len(filled)), key=lambda i: numbers[i] or 0.0, reverse=descending
        )
        filled = [filled[i] for i in order]
    else:
        filled.sort(key=lambda row: _text_sort_key(row[index]), reverse=descending)
    rows[:] = filled + blank


def query_rows(
    content: bytes,
    file_format: str,
    *,
    page: int = 1,
    page_size: int = 20,
    filter_column: str | None = None,
    filter_value: str | None = None,
    filter_operator: str = "exact",
    sort_column: str | None = None,
    sort_direction: str = "asc",
) -> dict[str, Any]:
    """Filter, sort and paginate a downloaded file. Returns a Tabular API-shaped dict."""
    rows_iter = (
        _iter_xlsx(content)
        if file_format == "xlsx"
        else _iter_csv(content, file_format)
    )
    try:
        header = next(row for row in rows_iter if any(cell.strip() for cell in row))
    except StopIteration:
        return {
            "data": [],
            "meta": {"total": 0, "page": page, "page_size": page_size},
            "links": {},
        }
    columns = _column_names(header)

    def column_index(name: str) -> int:
        if name not in columns:
            raise LocalQueryError(
                f"Column '{name}' does not exist in this resource. "
                f"Available columns: {', '.join(columns)}"
            )
        return columns.index(name)

    width = len(columns)
    matches: Callable[[list[str]], bool] = lambda row: True  # noqa: E731
    if filter_column and filter_value is not None:
        filter_index = column_index(filter_column)
        cell_matches = _build_filter(filter_operator, str(filter_value))
        matches = lambda row: cell_matches(row[filter_index])  # noqa: E731
    sort_index = column_index(sort_column) if sort_column else None

    start = (max(page, 1) - 1) * page_size
    end = start + page_size
    total = 0
    kept: list[list[str]] = []
    for raw in rows_iter:
        if not any(cell.strip() for cell in raw):
            continue
        row = (raw + [""] * width)[:width]
        if not matches(row):
            continue
        # Without a sort, only the requested page needs to stay in memory
        if sort_index is not None or start <= total < end:
            kept.append(row)
        total += 1

    if sort_index is not None:
        _sort_rows(kept, sort_index, descending=sort_direction == "desc")
        kept = kept[start:end]

    return {
        "data": [dict(zip(columns, row)) for row in kept],
        "meta": {"total": total, "page": max(page, 1), "page_size": page_size},
        "links": {"next": True} if end < total else {},
    }


async def fetch_resource_data(
    resource: dict[str, Any],
    **query: Any,
) -> dict[str, Any]:
    """
    Download a resource (from its catalog payload) and query it locally.

    Raises:
        LocalQueryError: unsupported format, file too large, or unknown column.
        niquests.HTTPError / BlockedAddressError: if the download fails.
    """
    resource_id = str(resource.get("id") or "")
    url = str(resource.get("url") or "")
    file_format = detect_format(resource)
    if not file_format or not url:
        raise LocalQueryError(
            f"Format '{resource.get('format') or 'unknown'}' cannot be queried "
            f"(supported: {', '.join(SUPPORTED_FORMATS)}). "
            "Fetch the resource URL directly instead."
        )
    filesize = resource.get("filesize")
    if isinstance(filesize, int) and filesize > max_bytes():
        raise LocalQueryError(
            f"The file ({format_size(filesize)}) is larger than the "
            f"{format_size(max_bytes())} limit for local queries. "
            "Fetch the resource URL directly instead."
        )
    content = await _download(resource_id, url, _resource_version(resource))
    return await asyncio.to_thread(query_rows, content, file_format, **query)
